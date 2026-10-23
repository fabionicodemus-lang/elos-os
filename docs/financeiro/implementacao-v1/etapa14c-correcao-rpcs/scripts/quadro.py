from pathlib import Path
import hashlib
import json
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
E = ROOT / 'evidencias'
SOURCE = ROOT.parent / 'etapa14b-exposicao/Elos_OS_Quadro_Etapa14B_Exposicao.xlsx'
OUT = ROOT / 'Elos_OS_Quadro_Etapa14C_Correcao_RPCs.xlsx'
old = openpyxl.load_workbook(SOURCE)
book = openpyxl.load_workbook(SOURCE)
notes = {
    'T42': 'Etapa14C: HTML independente confirma3 funções vulneráveis no banco vinculado à produção; HTTP ainda não comprovado. Candidata fecha EXECUTE direto e valida identidade/empresa/capacidade. Aplicação produtiva não realizada.',
    'T44': 'Etapa14C:33/33 controles,11 capacidades mínimas; regressões25/24/16/8/7/9/4/2,81 e153 conformes nos recortes. Fórmulas e30 cenários preservados;0 aceite formal.',
    'T45': 'Etapa14C: candidata para revisão em branch própria; produção continua bloqueada até preflight/staging e autorização explícita. Não aplicar a fila histórica; manter pendências globais da14/14B.'
}
allowed = set()
for row in book['Acompanhamento']:
    if row[0].value not in notes:
        continue
    for col, text in [('I', 'etapa14c-correcao-rpcs/README.md; TESTES.md; IMPLANTACAO-REVERSAO.md; evidencias/resultado.json'), ('S', notes[row[0].value])]:
        cell = f'{col}{row[0].row}'
        prior = book['Acompanhamento'][cell].value
        book['Acompanhamento'][cell] = (str(prior) + ' | ' if prior else '') + text
        allowed.add(('Acompanhamento', cell))
s = book.create_sheet('Etapa14C_Correcao')
s.append(['Função', 'Banco vinculado à produção', 'Exposição HTTP produtiva', 'Correção candidata', 'Aplicação', 'Aceite'])
for name in ['refresh_execution_service_contract_value', 'refresh_execution_service_contract_measurements', 'refresh_execution_work_order_values']:
    s.append([name, 'Vulnerável conforme HTML independente recebido', 'NÃO COMPROVADA', 'EXECUTE interno + guard usuário/empresa/capacidade; fórmula preservada', 'Não aplicada', 'Revisão/autorização pendentes; sem aceite financeiro'])
s = book.create_sheet('Etapa14C_Testes')
s.append(['Controle', 'Resultado', 'Obtido/evidência', 'Ambiente', 'Aceite financeiro'])
for r in json.loads((E / 'resultado.json').read_text())['checks']:
    s.append([r['name'], 'Conforme' if r['passed'] else 'Reprovado', json.dumps(r.get('actual', r.get('error')), ensure_ascii=False), 'Isolado fictício', 'Não registrado pelo executor'])
s = book.create_sheet('Etapa14C_Chamadores')
s.append(['Função', 'Assinatura', 'Arquivo', 'Linha', 'Permissões existentes', 'Chamadas', 'Grau estático', 'Limite'])
for r in json.loads((E / 'chamadores.json').read_text())['callers']:
    s.append([r['function'], r['signature'], r['path'], r['line'], ', '.join(r['permission_keys']), ', '.join(r['calls']), r['graph_depth'], 'Conferir definição/owner/ACL implantados; SQL dinâmico não certificado'])
s = book.create_sheet('Etapa14C_Implantacao')
s.append(['Passo', 'Estado', 'Responsável', 'Critério/evidência'])
for r in [
    ['Recuperar correção/local e base main auditada', 'Verificado', 'Executor técnico', '39d858c/10ae3f7/6d9fecb; merge isolado3eae06b de460147e'],
    ['Testar candidata e preservar fórmulas', 'Conforme no recorte', 'Executor técnico', '33/33 e regressões; relatórioTESTES.md'],
    ['Publicar branch para revisão', 'Conferir publicacao.json e commit remoto', 'Executor técnico', 'Não altera main nem aplica migração'],
    ['Preflight de produção e demais wrappers em staging', 'Pendente', 'Tecnologia/DBA a indicar', 'Metadados/owners/ACL, DataAPI, contextos de serviço e fluxos não executados'],
    ['Aprovar migração/janela/reversão', 'Pendente', 'Responsável autorizado da produção', 'Autorização explícita; revisão de segurança; não é aceite financeiro global'],
    ['Implantar correção', 'NÃO EXECUTADO', 'Operador autorizado a indicar', 'Somente arquivo14C aprovado; nenhuma fila/backfill/importação'],
]:
    s.append(r)
for name in ['Etapa14C_Correcao', 'Etapa14C_Testes', 'Etapa14C_Chamadores', 'Etapa14C_Implantacao']:
    book[name].freeze_panes = 'A2'
    book[name].auto_filter.ref = book[name].dimensions
    for col in book[name].columns:
        book[name].column_dimensions[col[0].column_letter].width = 46
book.calculation.fullCalcOnLoad = True
book.calculation.forceFullCalc = True
book.save(OUT)
after = openpyxl.load_workbook(OUT)
changed = []
for sheet in old:
    for row in sheet:
        for cell in row:
            if cell.value != after[sheet.title][cell.coordinate].value:
                assert (sheet.title, cell.coordinate) in allowed
                changed.append(f'{sheet.title}!{cell.coordinate}')
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
(E / 'verificacao-quadro.json').write_text(json.dumps(dict(
    source_sha256=sha(SOURCE), derived_sha256=sha(OUT), changed_cells=changed,
    all_30_case_values_expected_classifications_and_acceptances_unchanged=True,
    prior_task_statuses_formulas_and_other_cells_unchanged=True,
    formula_caches_recomputed=False), ensure_ascii=False, indent=2) + '\n')
