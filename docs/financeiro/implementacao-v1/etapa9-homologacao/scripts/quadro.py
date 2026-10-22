"""Cópia derivada da Etapa8; mantém aceites, regras e expected intactos."""
from pathlib import Path
import openpyxl,json,hashlib
from datetime import datetime
from zoneinfo import ZoneInfo
ROOT=Path(__file__).resolve().parents[1];E=ROOT/'evidencias';SOURCE=ROOT.parent/'etapa8-fluxo/Elos_OS_Quadro_Etapa8_Fluxo.xlsx';OUT=ROOT/'Elos_OS_Quadro_Etapa9_Homologacao.xlsx'
w=openpyxl.load_workbook(SOURCE);old=openpyxl.load_workbook(SOURCE);allowed=set();matrix=json.loads((E/'matriz-30.json').read_text())
notes={
'T19':('Em validação','Extratos fixos independentes CT03/13/14/23/22/06/10:9 posições fecham no centavo; banco real/corte comum ausentes. Conciliação não cria dinheiro. CT03 custo ainda parcial.'),
'T21':('Em andamento','CT06 conserva100.000 nos componentes; falta geração nativa pai/filhos, cobertura/custo e liberação condicionada. Não homologado integralmente.'),
'T24':('Bloqueada','CT10 executado e reprovado: cancelamento não cria direito10.000; estorno com devolução6.000 reabre AP10.000. Propor conversão documentada em restituição, sem alterar produção.'),
'T17':('Em andamento','CT11 oráculo60/40 e maior resto passa; multiobra continua pendente de ledger e teste integral. Quarentena multi_project_requires_ledger é preservada.'),
'T32':('Em validação','CT13/14: todos os campos fictícios passam com títulos, parciais, encargo, tarifa e extratos independentes. Dados reais/Auth/aceite nominal pendentes.'),
'T33':('Em validação','CT14 permuta liquida100.000 com banco30.000+compensação70.000; sem caixa fictício. CT22 aplicação passa, mas I18 e receita econômica não demonstrados.'),
'T34':('Em andamento','CT13 coorte30%, vencido6.120, aberto86.120; CT22 futuro contratual60.000/I1880.000 ainda sem mapa operacional; não presumir valores.'),
'T42':('Em andamento','Controles SQL preservados; cenário integral Auth/anexos/cache/exportações/IA/jobs ainda pendente.'),
'T44':('Em andamento','Etapa9:13 casos com cobertura técnica declarada (3 extrato independente novo;10 herdados),16 parciais e CT10 reprovado. Sete oráculos puros adicionais não promovem casos. Sem aceite financeiro; ponte real pendente.'),
'T45':('Bloqueada','Nenhuma liberação de produção. CT10, parciais, divergência histórica149.589,33, fontes/cortes e aceites nominais pendentes.'),
}
def put(sh,cell,value):w[sh][cell]=value;allowed.add((sh,cell))
for row in w['Acompanhamento']:
 if row[0].value not in notes:continue
 state,note=notes[row[0].value];n=row[0].row
 for col,value in [('F',state),('I','etapa9-homologacao/README.md; MATRIZ-30.md; PARCIAIS-18.md; evidencias/ponte.json'),('S',note)]:put('Acompanhamento',f'{col}{n}',value)
actual={r['id']:r['actual'] for r in matrix if r['id'] in ['CT13','CT14','CT23']};context=json.loads((E/'banco.json').read_text())['execution_context'];when=datetime.fromisoformat(context['executed_at'].replace('Z','+00:00')).astimezone(ZoneInfo('America/Sao_Paulo')).replace(tzinfo=None)
count=0
for row in w['Testes']:
 if row[0].value not in actual:continue
 val=actual[row[0].value]
 for key in row[1].value.split('.'):val=val[key]
 if isinstance(val,str) and row[1].value.endswith('_percent'):val=float(val)
 for col,value in [('G',val),('J','etapa9-homologacao/evidencias/comparacao-novos-integrais.json; banco.json; extratos-referencia.json — isolamento; sem aceite financeiro'),('K','Codex — observação automatizada isolada, não responsável financeiro'),('L',when)]:put('Testes',f'{col}{row[0].row}',value)
 count+=1
s=w.create_sheet('Etapa9_Confronto');s.append(['Caso','Regra','Classificação Etapa8','Classificação Etapa9','Esperado fixo (centavos)','Obtido (centavos/null)','Observado','Pendência específica','Evidência','Aceite financeiro'])
for r in matrix:s.append([r['id'],', '.join(r['rule_refs']),r['stage8_classification'],r['classification'],json.dumps(r['expected'],ensure_ascii=False),json.dumps(r['actual'],ensure_ascii=False),r['observed'],r['remaining_evidence'],'; '.join(r['evidence']),r['financial_acceptance']])
s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
for c in ['B','C','D']:s.column_dimensions[c].width=45
for c in ['E','F','G','H','I','J']:s.column_dimensions[c].width=90
w.calculation.fullCalcOnLoad=True;w.calculation.forceFullCalc=True;w.save(OUT);after=openpyxl.load_workbook(OUT);changes=[];same=0
for sh in old:
 for row in sh:
  for cell in row:
   if cell.value!=after[sh.title][cell.coordinate].value:
    assert (sh.title,cell.coordinate) in allowed;changes.append(f'{sh.title}!{cell.coordinate}')
   else:same+=1
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
(E/'verificacao-quadro.json').write_text(json.dumps(dict(source_sha256=sha(SOURCE),derived_sha256=sha(OUT),modified_existing_cells=changes,unchanged_cells=same,executed_test_rows=count,expected_formulas_acceptances_dependencies_unchanged=True,original_workbook_preserved=True,formula_caches_recomputed=False,added_sheet='Etapa9_Confronto'),ensure_ascii=False,indent=2)+'\n')
print(f'{count} linhas de observação atualizadas; campos de aceite preservados.')
