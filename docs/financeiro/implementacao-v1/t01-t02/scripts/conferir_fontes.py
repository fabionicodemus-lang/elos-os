"""T01/T02: leitura local de fontes preservadas; não conecta a serviços."""
import argparse
import csv
import hashlib
import json
from collections import Counter, defaultdict
from decimal import Decimal
from pathlib import Path

import openpyxl


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--package", required=True, type=Path)
    parser.add_argument("--sources", required=True, type=Path)
    args = parser.parse_args()
    package, sources = args.package, args.sources
    manifest = json.loads((package / "manifesto.json").read_text())
    mismatches = [name for name, digest in manifest["files_sha256"].items()
                  if sha(package / name) != digest]
    result = {"scope": "Fontes históricas locais; não é snapshot do Elos atual",
              "package_files_checked": len(manifest["files_sha256"]),
              "package_hash_mismatches": mismatches,
              "normative_sha256": sha(package / "regras/Elos_OS_Regras_Financeiras_v0_1.docx"),
              "historical_workbooks": {}, "source_file_hashes": {},
              "package_file_hashes": {}, "duplicates": [], "limitations": []}
    for root, key in [(sources, "source_file_hashes"), (package, "package_file_hashes")]:
        result[key] = {str(p.relative_to(root)): sha(p)
                       for p in sorted(root.rglob("*")) if p.is_file()}
    hashes = defaultdict(list)
    for p in (sources / "01_originais").rglob("*"):
        if p.is_file():
            hashes[sha(p)].append(str(p.relative_to(sources)))
    result["duplicates"] = [paths for paths in hashes.values() if len(paths) > 1]
    for source_id in ["F011", "F012", "F013", "F014", "F015", "F016"]:
        path = next((sources / "01_originais/obrigacoes" / source_id).glob("*.xlsx"))
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = book["Titulos_Koper"]
        rows = [row for row in sheet.iter_rows(min_row=5, values_only=True)
                if row[0] is not None]
        entry = {"file": str(path.relative_to(sources)), "sha256": sha(path),
                 "sheet": "Titulos_Koper", "header_row": 4,
                 "count": len(rows), "unique_external_ids": len({str(row[0]) for row in rows}),
                 "total_cents": int(sum(Decimal(str(row[4] or 0)) for row in rows) * 100),
                 "statuses": dict(Counter(row[6] for row in rows)),
                 "title_16032_occurrences": sum(str(row[0]) == "16032" for row in rows)}
        if source_id == "F011":
            missing_center = {str(row[0]) for row in rows
                              if not row[11] or row[11] == "Não identificado no relatório de fluxo"}
            missing_eap = {str(row[0]) for row in rows if not row[19]}
            entry.update({"missing_financial_center": len(missing_center),
                          "missing_eap": len(missing_eap),
                          "missing_both": len(missing_center & missing_eap),
                          "missing_either": len(missing_center | missing_eap),
                          "multiple_services_as_text": sum("|" in str(row[19] or "") for row in rows),
                          "columns": list(next(sheet.iter_rows(min_row=4, max_row=4, values_only=True)))})
        result["historical_workbooks"][source_id] = entry
        book.close()
    result["derived_controls"] = []
    with (sources / "controles_de_linhas.csv").open(encoding="utf-8-sig") as handle:
        for row in csv.DictReader(handle, delimiter=";"):
            result["derived_controls"].append(row)
    path = next((sources / "01_originais/orcamento/F017").glob("*.xlsx"))
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    errors = Counter()
    for sheet in book:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "e":
                    errors[str(cell.value)] += 1
    result["rev03_cached_errors"] = dict(errors)
    book.close()
    result["limitations"] = [
        "Título 16032 ausente destes seis arquivos: não demonstra cancelamento nem estado atual.",
        "Controle derivado é evidência documental; não substitui conciliação ao ERP/extrato.",
        "Hash comprova identidade do arquivo, não aprovação financeira nem completude.",
        "Não foram consultados dados atuais do Elos ou Koper; lista histórica dos 26 extras ausente."]
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if mismatches:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
