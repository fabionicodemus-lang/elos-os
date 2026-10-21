"""Regenera cópia do quadro; não preenche aceites nem altera oráculos/fórmulas."""
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
import hashlib
import json
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT.parent / 'base-validacao/Elos_OS_Quadro_Validacao_Base_T03_T08.xlsx'
OUT = ROOT / 'Elos_OS_Quadro_Etapa6_Koper.xlsx'
EVIDENCE = ROOT / 'evidencias'
workbook = openpyxl.load_workbook(SOURCE)
original = openpyxl.load_workbook(SOURCE)
notes = {
 'T09': ('Em validação', 'Manifestos/originais/hash e paginação executados; transporte real e amostra/aceite técnico pendentes.'),
 'T10': ('Em validação', 'Serial/concorrente/retry/revisão suportada passam no recorte AP; demais objetos e integração remota pendentes.'),
 'T11': ('Em validação', 'CT01/CT26 compatíveis no banco; ponte por ID/subtotal/falha. Falta posição real reconciliável e bruto/líquido.'),
 'T12': ('Bloqueada', 'Sem lista dos 26 extras e eventos atuais do 16032. Nenhuma exclusão ou conclusão histórica presumida.'),
 'T13': ('Bloqueada', 'Sem amostra/mapeamento histórico aprovado e efeito anterior comprovado. Contrapartes não saneadas.'),
 'T14': ('Em andamento', 'Execução/efeito local comprovados; configuração, run/deploy/efeito remotos e D12 não comprovados.'),
 'T15': ('Em andamento', 'Quarentena preserva entidade/obra/valores; nenhum fallback Flow. CT12 integral depende de parcial/caixa e D01.'),
 'T16': ('Em andamento', 'Vínculo explícito preserva título/identidade; faltam N:M e componentes/cadeia documental integral.'),
 'T17': ('Em andamento', 'A Apropriar e valores monetários/versionados preservados no recorte; D05 real e ledger multiobra/caixa pendentes.'),
}
allowed = set()
def put(sheet, coordinate, value):
    workbook[sheet][coordinate] = value
    allowed.add((sheet, coordinate))

for row in workbook['Acompanhamento']:
    task, number = row[0].value, row[0].row
    if task in notes:
        state, note = notes[task]
        put('Acompanhamento', f'F{number}', state)
        put('Acompanhamento', f'I{number}', 'docs/financeiro/implementacao-v1/etapa6-koper/README.md; evidencias/integracao.json; evidencias/comparacao-CT01-CT26.json; ACOMPANHAMENTO.md')
        put('Acompanhamento', f'S{number}', note)
    if task == 'T18':
        put('Acompanhamento', f'S{number}', 'Etapa 7: CT07 permanece crítico (parcial encerra título). Integração segura declaração sem criar quitação; ledger/componentes e sequência completa pendentes.')

model = json.loads((EVIDENCE / 'casos-obtidos.json').read_text())
executed = {c['id']: c['actual'] for c in model['cases'] if c['status'] == 'executado'}
assert set(executed) == {'CT01', 'CT26'}, 'Revisar o adaptador se a cobertura mudar'
when = datetime.fromisoformat(model['execution_context']['executed_at'].replace('Z', '+00:00')).astimezone(ZoneInfo('America/Sao_Paulo')).replace(tzinfo=None)
rows_checked = 0
for row in workbook['Testes']:
    if row[0].value not in executed:
        continue
    value = executed[row[0].value]
    for part in row[1].value.split('.'):
        value = value[part]
    if value is None:
        value = 'null'
    elif row[1].value.endswith('_percent'):
        value = float(value)
    number = row[0].row
    for column, val in [('G', value), ('J', 'etapa6-koper/evidencias/integracao.json; comparacao-CT01-CT26.json — banco/importação, sem aceite financeiro'), ('K', 'Codex — execução automatizada isolada'), ('L', when)]:
        put('Testes', f'{column}{number}', val)
    rows_checked += 1
workbook.calculation.fullCalcOnLoad = True
workbook.calculation.forceFullCalc = True
workbook.save(OUT)
after = openpyxl.load_workbook(OUT)
changed, unchanged = [], 0
for sheet in original:
    for row in sheet:
        for cell in row:
            if cell.value != after[sheet.title][cell.coordinate].value:
                assert (sheet.title, cell.coordinate) in allowed
                changed.append(f'{sheet.title}!{cell.coordinate}')
            else:
                unchanged += 1
assert len(changed) == len(allowed)
digest = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
(EVIDENCE / 'verificacao-quadro.json').write_text(json.dumps({
 'source_sha256': digest(SOURCE), 'derived_sha256': digest(OUT),
 'modified_cells': changed, 'unchanged_cells_checked': unchanged,
 'executed_test_rows': rows_checked,
 'expected_formulas_dependencies_acceptances_unchanged': True,
 'formula_values_recomputed': False, 'execution_timezone': 'America/Sao_Paulo',
}, ensure_ascii=False, indent=2) + '\n')
print(f'{len(changed)} células de entrega/resultados; {unchanged} preservadas; aceites intactos.')
