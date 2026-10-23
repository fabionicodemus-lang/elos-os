from pathlib import Path
import hashlib
import json
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
E = ROOT / 'evidencias'
SOURCE = ROOT.parent / 'etapa14-isolamento/Elos_OS_Quadro_Etapa14_Isolamento.xlsx'
OUT = ROOT / 'Elos_OS_Quadro_Etapa14B_Exposicao.xlsx'
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
old = openpyxl.load_workbook(SOURCE)
book = openpyxl.load_workbook(SOURCE)
allowed = set()
notes = {
    'T42': 'Etapa14B: exposição produtiva NÃO VERIFICADA; sem conexão/identidade de metadados ou atestado da implantação. 3 candidatas locais presentes; coletor e contenção preparados, sem aplicação produtiva. Canais restantes NÃO CERTIFICADOS.',
    'T44': 'Etapa14B: 7 controles do coletor/artefatos,8 guards e4 OS;81 regressões e142 testes padrão aprovados. Não reclassifica os30 casos;0 aceite formal. HTTP342 é evidência herdada da14.',
    'T45': 'Etapa14B: manter BLOQUEIO global; NÃO FOI POSSÍVEL COMPROVAR versão em produção. Exigir snapshot atestado/configuração de API, revisão de dependências e autorização explícita antes de qualquer aplicação.'
}
for row in book['Acompanhamento']:
    task = row[0].value
    if task not in notes:
        continue
    for col, text in [('I', 'etapa14b-exposicao/README.md; REMEDIACAO.md; evidencias/testes-metadados.json'), ('S', notes[task])]:
        cell = f'{col}{row[0].row}'
        prior = book['Acompanhamento'][cell].value
        book['Acompanhamento'][cell] = (str(prior) + ' | ' if prior else '') + text
        allowed.add(('Acompanhamento', cell))

s = book.create_sheet('Etapa14B_Exposicao')
s.append(['Falha', 'RPC', 'Ambiente comprovado', 'Versão implantada', 'Caminho potencial', 'Impacto', 'Criticidade potencial', 'Certeza', 'Evidência', 'Aceite financeiro'])
for r in json.loads((E / 'riscos.json').read_text()):
    s.append([r['id'], r['signature'], r['proven_environment'], r['production_version_status'], r['potential_path'], r['impact'], r['severity_if_exposed'], r['certainty'], r['evidence'], 'Não registrado pelo executor'])
s = book.create_sheet('Etapa14B_Plano')
s.append(['Passo', 'Estado', 'Dependência/critério', 'Responsável', 'Evidência', 'Autorização'])
for row in [
    ['Identificar implantação e coletar metadados', 'Pendente de acesso', 'Conexão ou snapshot atestado; não ler dados nem chamar RPC', 'Tecnologia/operador a indicar', 'sql/inventario-metadados.sql', 'Leitura autorizada nesta tarefa; acesso ausente'],
    ['Comparar definição, ACL, helpers e gateway', 'Preparado; produção não verificada', 'Snapshot íntegro e configurações externas não secretas', 'Segurança/DBA a indicar', 'scripts/comparar.py; README.md', 'Sem aceite produtivo'],
    ['Validar candidatas locais', 'Conforme no recorte', '7/7,8/8,4/4;81/81;142/142. DDL remoto desconhecido', 'Executor técnico', 'evidencias/execucoes.json', 'Não é homologação financeira'],
    ['Reproduzir DDL implantado em staging', 'Pendente', 'Massa fictícia, integração/HTTP e callers; nenhuma cópia de dados reais', 'Tecnologia/segurança', 'REMEDIACAO.md', 'Acesso/DDL ausente'],
    ['Aplicar correções condicionais', 'Não executado; não autorizado em produção', 'Falha implantada comprovada; revisão; matriz nominal; janela', 'Operador de produção a indicar', 'Duas migrações existentes da14', 'Exige autorização explícita futura'],
    ['Contenção/reversão segura', 'Artefato local testado', 'Não restaurar PUBLIC/anon nem corpo vulnerável; revisar grants alternativos', 'Operação/segurança a indicar', 'sql/contencao-condicional.sql', 'Exige autorização explícita futura'],
]:
    s.append(row)
s = book.create_sheet('Etapa14B_Canais')
s.append(['Canal', 'Evidência herdada', 'Lacuna', 'Estado'])
for line in (ROOT / 'CANAIS-PENDENTES.md').read_text().splitlines():
    if line.startswith('|') and not line.startswith(('|---', '| Canal')):
        s.append([c.strip().replace('**', '') for c in line.strip('|').split('|')])
for name in ['Etapa14B_Exposicao', 'Etapa14B_Plano', 'Etapa14B_Canais']:
    book[name].freeze_panes = 'A2'
    book[name].auto_filter.ref = book[name].dimensions
    for col in book[name].columns:
        book[name].column_dimensions[col[0].column_letter].width = 44
book.calculation.fullCalcOnLoad = True
book.calculation.forceFullCalc = True
book.save(OUT)
after = openpyxl.load_workbook(OUT)
modified = []
for sh in old:
    for row in sh:
        for cell in row:
            if cell.value != after[sh.title][cell.coordinate].value:
                assert (sh.title, cell.coordinate) in allowed, (sh.title, cell.coordinate)
                modified.append(f'{sh.title}!{cell.coordinate}')
(E / 'verificacao-quadro.json').write_text(json.dumps({
    'source_sha256': sha(SOURCE), 'derived_sha256': sha(OUT), 'modified_cells': modified,
    'all_prior_sheet_cells_identical_except_explicit_evidence_and_notes': True,
    'all_30_case_expected_actual_and_classifications_unchanged': True,
    'formal_acceptance_task_statuses_and_formulas_unchanged': True,
    'formula_caches_recomputed': False
}, ensure_ascii=False, indent=2) + '\n')
