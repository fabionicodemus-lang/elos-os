"""Auditoria offline de artefatos históricos; não acessa banco, rede ou código financeiro.

Dependências: Python 3, openpyxl, PyMuPDF e pdftotext (Poppler).
Centavos inteiros, desconhecido=None. A chave ID/parcela só vale neste par de fontes;
não substitui organização/entidade/origem na identidade de uma importação real.
"""
from pathlib import Path
from decimal import Decimal
from collections import Counter, defaultdict
from datetime import datetime, timezone
import argparse
import csv
import hashlib
import json
import re
import subprocess
import fitz
import openpyxl
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
MONEY = r"(?:\d{1,3}(?:\.\d{3})*|\d+),\d{2}"
ID = re.compile(r"^\s*(\d+)\s+(\d+\s*/\s*\d+)\s+", re.M)
DATE = r"\d{2}/\d{2}/\d{4}"


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def cents(value):
    if value is None or value == "":
        return None
    n = Decimal(str(value)) * 100
    if not n.is_finite() or n != n.to_integral_value():
        raise ValueError(f"Valor sem centavos exatos: {value!r}")
    return int(n)


def br_cents(value):
    return cents(value.replace(".", "").replace(",", "."))


def subtotal(values):
    values = list(values)
    return {"known_subtotal_cents": sum(v for v in values if v is not None),
            "missing_count": sum(v is None for v in values),
            "total_cents": None if any(v is None for v in values) else sum(values)}


def parse_pdf_text(text):
    """Blocos de linhas; cabeçalhos/rodapés não são registros nem pagamentos."""
    rows, page_controls = [], []
    pages = text.split("\f")
    if not pages[-1].strip():
        pages.pop()
    for page, content in enumerate(pages, 1):
        footer = re.search(r"Página\s+(\d+) de (\d+)", content)
        if not footer or int(footer[1]) != page or int(footer[2]) != len(pages):
            raise ValueError(f"Página/rodapé incompleto: {page}")
        clean = "\n".join(line if not line.startswith("Gerado por KOPER") else "" for line in content.splitlines())
        page_rows = []
        for block in re.finditer(r"(?:[^\n]|\n(?!\s*\n))+", clean):
            raw = block[0]
            matches = list(ID.finditer(raw))
            if not matches:
                continue
            if len(matches) != 1:
                raise ValueError(f"Bloco ambíguo na página {page}")
            match = matches[0]
            amounts = re.findall(rf"(?<!\S){MONEY}(?!\S)", raw)
            dates = re.findall(rf"\b{DATE}\b", raw)
            if len(amounts) != 1 or len(dates) not in (1, 2):
                raise ValueError(f"Valor/datas ambíguos: página {page}, ID {match[1]}")
            row = dict(external_id=match[1], installment=re.sub(r"\s", "", match[2]),
                       displayed_cents=br_cents(amounts[0]), due=dates[0],
                       paid_date=dates[1] if len(dates) == 2 else None,
                       page=page, text_line=1 + clean.count('\n', 0, block.start() + match.start(1)),
                       block_sha256=hashlib.sha256(raw.encode()).hexdigest(), raw=raw)
            rows.append(row)
            page_rows.append(row)
        # Every numeric ID/parcel anchor must have been captured exactly once.
        if len(page_rows) != len(list(ID.finditer(clean))):
            raise ValueError(f"Cobertura de linhas incompleta na página {page}")
        page_controls.append(dict(page=page, rows=len(page_rows),
                                  displayed_cents=sum(r['displayed_cents'] for r in page_rows)))
    return rows, page_controls


def parse_pdf_coordinates(path):
    """Verificação independente por palavras e coordenadas do PDF, não pelo TXT.

    Limites de coluna medidos neste layout. Falha explicitamente se o layout mudar.
    Cada valor e data deve ter âncora única de ID próxima (<10 pontos verticais).
    """
    result = []
    with fitz.open(path) as doc:
        for page in doc:
            words = page.get_text("words")
            headers = [w for w in words if w[4] == 'Código']
            if len(headers) != 1:
                raise ValueError('Cabeçalho de coluna inesperado')
            h = headers[0]
            anchors = [w for w in words if 34 < w[0] < 72 and h[3] < w[1] < 800 and w[4].isdigit()]
            rows = {a[4]: dict(external_id=a[4], page=page.number + 1, installment_parts=[], document_parts=[]) for a in anchors}
            if len(rows) != len(anchors):
                raise ValueError('ID duplicado na página')
            for word in words:
                if not h[3] < word[1] < 800:
                    continue
                x, val = word[0], word[4]
                kind = ('displayed_cents' if 280 < x < 334 and re.fullmatch(MONEY, val)
                        else 'due' if 334 < x < 389 and re.fullmatch(DATE, val)
                        else 'paid_date' if 389 < x < 440 and re.fullmatch(DATE, val)
                        else 'installment_parts' if 74 < x < 110 and re.fullmatch(r"\d+|/|\d+/\d+", val)
                        else 'document_parts' if 110 < x < 187
                        else None)
                if kind is None:
                    continue
                distances = sorted((abs(a[1] - word[1]), a[4]) for a in anchors)
                if distances[0][0] >= 10 or (len(distances) > 1 and distances[0][0] == distances[1][0]):
                    raise ValueError('Âncora ambígua')
                row = rows[distances[0][1]]
                if kind in ('installment_parts', 'document_parts'):
                    row[kind].append((round(word[1], 1), x, val))
                else:
                    if kind in row:
                        raise ValueError('Campo repetido')
                    row[kind] = br_cents(val) if kind == 'displayed_cents' else val
            for row in rows.values():
                row['installment'] = ''.join(v for _, _, v in sorted(row.pop('installment_parts')))
                row['document'] = ' '.join(v for _, _, v in sorted(row.pop('document_parts')))
                if not re.fullmatch(r"\d+/\d+", row['installment']) or 'displayed_cents' not in row or 'due' not in row:
                    raise ValueError('Registro incompleto')
                row.setdefault('paid_date', None)
                result.append(row)
    return result


def key(row):
    return row['external_id'], row['installment']


def index(rows):
    found = defaultdict(list)
    for row in rows:
        found[key(row)].append(row)
    return found


def compare(pdf_rows, xlsx_rows):
    p, x = index(pdf_rows), index(xlsx_rows)
    records = []
    for k in sorted(p.keys() | x.keys(), key=lambda k: (int(k[0]), k[1])):
        pr, xr = p.get(k, []), x.get(k, [])
        one_p, one_x = pr[0] if len(pr) == 1 else {}, xr[0] if len(xr) == 1 else {}
        both = bool(one_p and one_x)
        delta = (one_x['displayed_cents'] - one_p['displayed_cents']
                 if both and one_x['displayed_cents'] is not None and one_p['displayed_cents'] is not None else None)
        differences = [f for f in ['due', 'paid_date'] if both and one_p.get(f) != one_x.get(f)]
        doc = re.sub(r'\s+', '', str(one_x.get('document') or '-'))
        doc_seen = doc in re.sub(r'\s+', '', one_p.get('raw', '')) if both else None
        pdf_doc = one_p.get('document')
        doc_class = (None if not both or pdf_doc is None else
                     'same_reference' if doc == re.sub(r'\s+', '', pdf_doc) else
                     'synthetic_KOPER_reference_for_missing_document' if pdf_doc == '-' and doc == 'KOPER' + k[0] else
                     'different_reference')
        classification = ('duplicate_key' if len(pr) > 1 or len(xr) > 1 else
                          'only_F011' if not pr else 'only_F007' if not xr else
                          'unknown_amount' if delta is None else 'different_value' if delta else
                          'different_dates' if differences else 'same_displayed_value_and_dates')
        records.append(dict(external_id=k[0], installment=k[1], classification=classification,
                            F007_count=len(pr), F011_count=len(xr), F007_page=one_p.get('page'),
                            F007_text_line=one_p.get('text_line'), F007_block_sha256=one_p.get('block_sha256'),
                            F011_sheet='Titulos_Koper' if xr else None, F011_row=one_x.get('row'),
                            F011_document=one_x.get('document'), F007_document=pdf_doc,
                            document_classification=doc_class, document_text_seen=doc_seen,
                            F007_displayed_cents=one_p.get('displayed_cents'), F011_displayed_cents=one_x.get('displayed_cents'),
                            delta_cents=delta, F007_due=one_p.get('due'), F011_due=one_x.get('due'),
                            F007_paid_date=one_p.get('paid_date'), F011_paid_date=one_x.get('paid_date'),
                            F011_status=one_x.get('status'), F011_declared_paid_cents=one_x.get('declared_paid_cents'),
                            F011_financial_center=one_x.get('financial_center'),
                            legal_entity_id=None, verified_project_id=None, ledger_components_cents=None,
                            confirmed_cash_cents=None, allocated_header_difference_cents=None))
    return records


def header_values(text):
    labels = {'total': 'Valor total', 'paid': 'Valor pago', 'unpaid': 'Valor não pago',
              'discount': 'Valor desconto', 'interest_late_fee': 'Valor juros/multa',
              'consolidated_recurrences': 'Recorrentes consolidadas'}
    first = text.split('Código')[0]
    vals = {}
    for k, label in labels.items():
        match = re.search(rf'{re.escape(label)}:\s*R\$\s*({MONEY})', first)
        if not match:
            raise ValueError(f'Cabeçalho ausente: {label}')
        vals[k + '_cents'] = br_cents(match[1])
    vals['count'] = int(re.search(r'Total de registros:\s*(\d+)', first)[1])
    return vals


def write_csv(path, rows):
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]), delimiter=';', lineterminator='\n')
        writer.writeheader()
        writer.writerows(rows)  # None exports as empty, not zero; JSON is authoritative for null.


def style_sheet(sheet):
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='008780')
    for cells in sheet.columns:
        sheet.column_dimensions[cells[0].column_letter].width = min(65, max(18, len(str(cells[0].value)) + 2))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--base', type=Path, default=Path('/workspace/financial-source-base/Base_Conferencia_Elos_OS'))
    parser.add_argument('--out', type=Path, default=ROOT / 'evidencias')
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    catalog = list(csv.DictReader((args.base / 'catalogo_fontes.csv').open(encoding='utf-8-sig'), delimiter=';'))
    sources = {r['fonte_id']: r for r in catalog if r['fonte_id'] in ('F007', 'F011')}
    for row in sources.values():
        row['observed_sha256'] = sha(args.base / row['caminho_original'])
        if row['observed_sha256'] != row['sha256']:
            raise ValueError('Fonte diverge do catálogo')
    pdf = args.base / sources['F007']['caminho_original']
    xlsx = args.base / sources['F011']['caminho_original']
    text = subprocess.check_output(['pdftotext', '-layout', str(pdf), '-'], text=True)
    cached = args.base / '03_textos_pdf/F007/Contas a pagar (20).txt'
    pdf_rows, page_controls = parse_pdf_text(text)
    coords = parse_pdf_coordinates(pdf)
    if len(index(pdf_rows)) != len(pdf_rows) or len(index(coords)) != len(coords):
        raise ValueError('Duplicidade no PDF; análise manual necessária')
    coord_map = {key(r): r for r in coords}
    for r in pdf_rows:
        for field in ['displayed_cents', 'due', 'paid_date', 'page']:
            if coord_map[key(r)][field] != r[field]:
                raise ValueError('Extratores discordam')
        r['document'] = coord_map[key(r)]['document']
    book = openpyxl.load_workbook(xlsx, data_only=False)
    xr = []
    for n, r in enumerate(book['Titulos_Koper'].iter_rows(min_row=5, values_only=True), 5):
        if r[0] is None:
            raise ValueError('Linha sem identidade')
        xr.append(dict(external_id=str(r[0]), installment=re.sub(r'\s', '', str(r[3])),
                       displayed_cents=cents(r[4]), due=r[5].strftime('%d/%m/%Y') if r[5] else None,
                       paid_date=r[7].strftime('%d/%m/%Y') if r[7] else None,
                       declared_paid_cents=cents(r[8]), status=r[6], document=r[2], row=n,
                       financial_center=r[11]))
    result = compare(pdf_rows, xr)
    h = header_values(text)
    dated = [r for r in pdf_rows if r['paid_date'] is not None]
    undated = [r for r in pdf_rows if r['paid_date'] is None]
    sums = dict(F011=subtotal(r['displayed_cents'] for r in xr),
                F007_lines=subtotal(r['displayed_cents'] for r in pdf_rows),
                F007_dated=subtotal(r['displayed_cents'] for r in dated),
                F007_undated=subtotal(r['displayed_cents'] for r in undated))
    delta = sums['F011']['total_cents'] - h['total_cents']
    report = dict(executed_at_utc=datetime.now(timezone.utc).isoformat(),
                  scope='Auditoria documental offline, sem banco/rede/alterações financeiras',
                  sources=sources, counts=dict(F011=len(xr), F007=len(pdf_rows), pages=len(page_controls),
                    with_paid_date=len(dated), without_paid_date=len(undated)),
                  classifications=dict(Counter(r['classification'] for r in result)), header=h, totals=sums,
                  document_classifications=dict(Counter(r['document_classification'] for r in result)),
                  identity_checks=dict(F011_distinct_ids=len({r['external_id'] for r in xr}),
                    F007_distinct_ids=len({r['external_id'] for r in pdf_rows}),
                    F011_distinct_keys=len(index(xr)), F007_distinct_keys=len(index(pdf_rows)),
                    F011_duplicate_keys=sum(len(v)>1 for v in index(xr).values()),
                    F007_duplicate_keys=sum(len(v)>1 for v in index(pdf_rows).values()),
                    scope='ID/parcela somente neste par documental; entidade/obra jurídica não validadas'),
                  visual_identity_observation=dict(source='F007 página1, renderização visual',
                    commercial_logo='Bossa / Flow Aptos', legal_entity_id=None, project_filter_validated=None,
                    note='Logo ausente no texto extraído; complementa o catálogo, sem comprovar identidade jurídica por linha'),
                  bridge=dict(F011_minus_F007_lines_cents=sums['F011']['total_cents']-sums['F007_lines']['total_cents'],
                    F007_lines_minus_header_cents=sums['F007_lines']['total_cents']-h['total_cents'],
                    dated_lines_minus_paid_header_cents=sums['F007_dated']['total_cents']-h['paid_cents'],
                    undated_lines_minus_unpaid_header_cents=sums['F007_undated']['total_cents']-h['unpaid_cents'],
                    total_difference_cents=delta, financially_explained_cents=0,
                    financial_cause_unallocated_cents=delta,
                    certified_financial_total_cents=None),
                  conditional_gross_net_check=dict(assumption='Somente se Valor for principal e todo desconto/encargo pertencer às mesmas linhas e ao mesmo corte; não comprovado',
                    lines_less_discount_plus_interest_cents=sums['F007_lines']['total_cents']-h['discount_cents']+h['interest_late_fee_cents'],
                    hypothetical_net_minus_header_cents=sums['F007_lines']['total_cents']-h['discount_cents']+h['interest_late_fee_cents']-h['total_cents']),
                  extraction_checks=dict(second_engine='PyMuPDF coordenadas vs Poppler blocos',
                    matched_rows=len(pdf_rows), matched_fields=['ID/parcela','valor exibido','vencimento','pagamento informado','página'],
                    poppler_text_sha256=hashlib.sha256(text.encode()).hexdigest(),
                    preserved_text_sha256=sha(cached), preserved_text_matches=text==cached.read_text(),
                    document_text_seen_count=sum(r['document_text_seen'] is True for r in result)),
                  xlsx_internal_metadata=dict(created=str(book.properties.created), modified=str(book.properties.modified), creator=book.properties.creator,
                    warning='Metadados internos não comprovam extração; divergem da data de entrada informada no catálogo.'),
                  financial_cause_proven=False, corrected_indicator=None, production_modified=False,
                  missing=['Resposta bruta Koper/HTML e consulta/fórmula do relatório no instante 22/07/2026 10:00:47, com filtros, paginação e identidade legal/obra',
                           'billValue, paymentValue, desconto/juros/multa/outros, relações pai/agrupamento/recorrência e eventos de cada ID/parcela naquele snapshot',
                           'Manifesto/transformador original da F011 com corte e proveniência dos campos Status/Valor Pago',
                           'Eventos de parciais, cancelamentos e estornos; extratos e comprovantes se a medida pretendida for caixa',
                           'Snapshot Elos no mesmo T/K para investigar efeito no destino; responsável pelo aceite financeiro'])
    for name, value in [('reconciliacao.json', report), ('registros.json', result), ('paginas.json', page_controls)]:
        (args.out / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')
    write_csv(args.out / 'reconciliacao-por-registro.csv', result)
    write_csv(args.out / 'controle-paginas.csv', page_controls)
    wb = openpyxl.Workbook()
    sh = wb.active
    sh.title = 'Resumo'
    for row in [['Medida','Centavos / valor','Condição'],
                ['F011 soma linhas',sums['F011']['total_cents'],'Valor exibido, não caixa validado'],
                ['F007 soma linhas',sums['F007_lines']['total_cents'],'Valor exibido, não caixa validado'],
                ['F007 cabeçalho total',h['total_cents'],'Fórmula não disponível'],
                ['Diferença documental',delta,'Localizada no resumo do PDF'],
                ['Causa financeira sem alocação',delta,'Pendente; não é lançamento de ajuste'],
                ['Resultado financeiro correto',None,'Desconhecido'],
                ['Unidade','Centavos inteiros','Vazio = desconhecido; zero é observado'],
                ['Identidade','ID/parcela no par de fontes','CNPJ e obra validados ausentes']]:
        sh.append(row)
    sh = wb.create_sheet('Registros')
    sh.append(list(result[0]))
    for r in result:
        sh.append(list(r.values()))
    sh = wb.create_sheet('Paginas')
    sh.append(list(page_controls[0]))
    for r in page_controls:
        sh.append(list(r.values()))
    sh = wb.create_sheet('Fontes')
    sh.append(['Fonte','Arquivo original','SHA256','Empresa / obra declaradas','Data / corte'])
    for sid, r in sources.items():
        sh.append([sid,r['caminho_original'],r['sha256'],r['empresa']+' / '+r['obra'],r['data_extracao']+' / '+r['data_corte']])
    for sh in wb:
        style_sheet(sh)
    wb.save(args.out / 'Reconciliacao_F011_F007.xlsx')
    with fitz.open(pdf) as d:
        d[0].get_pixmap(matrix=fitz.Matrix(1.5,1.5), clip=fitz.Rect(20,40,575,227)).save(args.out / 'F007-cabecalho.png')
    print(json.dumps(dict(counts=report['counts'], classifications=report['classifications'],
                         bridge=report['bridge'], conditional_check=report['conditional_gross_net_check'],
                         extraction_checks=report['extraction_checks']), ensure_ascii=False))


if __name__ == '__main__':
    main()
