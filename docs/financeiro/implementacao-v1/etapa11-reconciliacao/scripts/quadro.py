"""Cópia aditiva do quadro da Etapa10; não altera aceites/oráculos/estados CT."""
from pathlib import Path
import hashlib
import json
import openpyxl
from reconciliar import style_sheet

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT.parent / 'etapa10-ct10/Elos_OS_Quadro_Etapa10_CT10.xlsx'
OUT = ROOT / 'Elos_OS_Quadro_Etapa11_Reconciliacao.xlsx'
old = openpyxl.load_workbook(SOURCE)
book = openpyxl.load_workbook(SOURCE)
allowed = set()
notes = {
    'T11': 'Etapa11: 3261 IDs/parcelas, valores e datas idênticos F011/F007; 165 páginas conferidas. Diferença149.589,33 localizada entre linhas e resumo Pago do próprio PDF. Causa financeira/componentes e fórmula de origem ausentes; não concluída.',
    'T38': 'Etapa11: igualdade documental não certifica caixa, competência, obrigação ou cobertura legal. 724 referências KOPER são substitutos de documento ausente. Manter ressalvas e residual149.589,33 sem ajuste.',
    'T44': 'Etapa11: ponte documental por registro pronta para revisão; causa econômica do residual149.589,33 não comprovada. Nenhum estado dos30 casos ou aceite alterado. CT10 daEtapa10 preservado.',
    'T45': 'Etapa11 diagnóstica encerrada com fonte bruta/fórmula/eventos históricos pendentes. Produção inalterada. Não libera implantação nem homologação global.',
}
for row in book['Acompanhamento']:
    if row[0].value not in notes:
        continue
    n = row[0].row
    for col, addition in [('I', 'etapa11-reconciliacao/README.md; evidencias/reconciliacao.json; evidencias/Reconciliacao_F011_F007.xlsx'), ('S', notes[row[0].value])]:
        cell = book['Acompanhamento'][f'{col}{n}']
        cell.value = (cell.value + ' | ' if cell.value else '') + addition
        allowed.add(('Acompanhamento', cell.coordinate))

sheet = book.create_sheet('Etapa11_Reconciliacao')
for row in [
    ['Critério', 'Resultado', 'Evidência', 'Pendência / condição de aceite'],
    ['Confronto de registros', '3261/3261 iguais em ID/parcela/valor/vencimento/pagamento informado', 'etapa11-reconciliacao/evidencias/registros.json', 'Aceite técnico nominal não preenchido'],
    ['Documentos', '2537 referências iguais; 724 substitutos KOPER onde F007 mostra -', 'etapa11-reconciliacao/evidencias/Reconciliacao_F011_F007.xlsx', 'Substituto não comprova NF/contrato/vínculo jurídico'],
    ['Ponte de valores exibidos', 'F011=linhas F007=13.710.826,32; cabeçalho F007=13.561.236,99', 'etapa11-reconciliacao/evidencias/reconciliacao.json', 'Residual financeiro149.589,33 sem causa comprovada'],
    ['Localização do residual', '13.687.213,70 nas3255 linhas datadas menos13.537.624,37 no resumo Pago', 'etapa11-reconciliacao/README.md', 'Não alocar diferença a títulos por suposição'],
    ['Causa raiz financeira', 'Não verificável com os artefatos disponíveis', 'etapa11-reconciliacao/README.md — fontes faltantes', 'Exigir snapshot bruto/consulta/componentes/eventos no mesmoT/K'],
    ['Testes offline', '12 diagnóstico +28 financeiros +81 regressão +142 suíte geral aprovados', 'etapa11-reconciliacao/evidencias/*log', 'Não representa nova execução dos controles25/24 nem homologação dos30'],
    ['Estado T11/T44', 'Mantidos Em validação / Em andamento', 'Acompanhamento e verificação do quadro', 'Dependências e aceites anteriores preservados'],
    ['Correção', 'Nenhuma implementada', 'git diff e manifesto-entrega.json', 'Causa e aprovação necessárias antes de implementar'],
    ['Aceite financeiro', 'Pendente — não preenchido pelo executor', 'Todos os campos anteriores preservados', 'Responsável financeiro'],
    ['Produção/Koper', 'Sem conexão financeira remota ou escrita', 'manifesto-entrega.json', 'T45 permanece Bloqueada'],
]:
    sheet.append(row)
style_sheet(sheet)
for col in 'ABCD':
    sheet.column_dimensions[col].width = 65
book.calculation.fullCalcOnLoad = True
book.calculation.forceFullCalc = True
book.save(OUT)
after = openpyxl.load_workbook(OUT)
changed = []
for sh in old:
    for row in sh:
        for c in row:
            if c.value != after[sh.title][c.coordinate].value:
                assert (sh.title, c.coordinate) in allowed, (sh.title, c.coordinate)
                changed.append(f'{sh.title}!{c.coordinate}')
assert len(changed) == 8
sha = lambda path: hashlib.sha256(path.read_bytes()).hexdigest()
out = dict(source_sha256=sha(SOURCE), derived_sha256=sha(OUT), modified_existing_cells=changed,
           old_sheets_preserved=old.sheetnames, all_30_case_cells_unchanged=True,
           expected_formulas_acceptances_dependencies_and_statuses_unchanged=True,
           financial_acceptance_recorded=False, formula_caches_recomputed=False)
(ROOT/'evidencias/verificacao-quadro.json').write_text(json.dumps(out, ensure_ascii=False, indent=2)+'\n')
print(json.dumps(out, ensure_ascii=False))
